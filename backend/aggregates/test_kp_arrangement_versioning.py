"""Backward-compatible Key-Population (KP) arrangement handling.

The reported KP arrangement changed; some users still submit the old one. These
tests pin the guarantees:

  * KP is identified by its canonical *label* (via the embedded ``_cellmap``),
    never by column/row position — so old and new arrangements map to the same
    internal categories.
  * An older arrangement (older/absent version marker, or a label that resolved
    only via an alias) is detected and surfaced with a clear notice.
  * Unrecognised KP categories are flagged + rejected, never silently stored or
    guessed into another category.
  * Aliased/duplicate labels collapse to canonical, sum-preserving, and the merge
    is reported.
  * Historical aggregates are never rewritten by an import of other data.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from aggregates import reporting_workbook as rw
from aggregates.models import Aggregate
from indicators.disaggregation import (
    KP_ARRANGEMENT_VERSION,
    STANDARD_KP_VALUES,
    classify_kp_label,
)
from indicators.models import Indicator
from organizations.models import Organization
from projects.models import (
    Project,
    ProjectIndicator,
    ProjectIndicatorOrganizationTarget,
)
from users.models import User

KP_CONFIG = {
    "enabled": True,
    "layout": "nested-matrix",
    "dimensions": [
        {"key": "key_population", "label": "Key Population",
         "values": ["MSM", "FSW", "PWD", "PWID", "PWUD", "LGBTQI+", "GENERAL POP."]},
        {"key": "sex", "label": "Sex", "values": ["Male", "Female"]},
        {"key": "age_band", "label": "Age Range", "values": ["10-14", "15-19"]},
    ],
}


def make_workbook(cells, *, indicator_id, project_id, org_id,
                  kp_arrangement="v2", code="KP-1"):
    """Build a minimal valid SESIGO workbook placing each ``cells`` entry at its
    own coordinate, with a self-describing ``_cellmap``. ``cells`` is a list of
    ``(primary, secondary, band, value)``. ``kp_arrangement=None`` omits the
    marker (mimics a pre-feature workbook). Returns xlsx bytes."""
    wb = Workbook()
    form = wb.active
    form.title = rw.SHEET_FORM
    meta = wb.create_sheet(rw.SHEET_META)
    cellmap = wb.create_sheet(rw.SHEET_CELLMAP)

    meta.cell(row=1, column=1, value="key")
    meta.cell(row=1, column=2, value="value")
    pairs = [
        ("workbook_version", rw.WORKBOOK_VERSION),
        ("project_id", project_id),
        ("organization_id", org_id),
        ("quarter", 3),
        ("fiscal_start_year", 2025),
        ("period_start", "2025-10-01"),
        ("period_end", "2025-12-31"),
    ]
    if kp_arrangement is not None:
        pairs.insert(1, ("kp_arrangement_version", kp_arrangement))
    for i, (k, v) in enumerate(pairs, start=2):
        meta.cell(row=i, column=1, value=k)
        meta.cell(row=i, column=2, value=v)

    cellmap.append(["indicator_id", "indicator_code", "kind",
                    "primary", "secondary", "band", "coordinate"])
    r = 5
    for (primary, secondary, band, value) in cells:
        coord = f"D{r}"
        form[coord] = value
        cellmap.append([indicator_id, code, "input", primary, secondary, band,
                        f"'{rw.SHEET_FORM}'!{coord}"])
        r += 1

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class KpClassifyTest(TestCase):
    """Unit tests for the catalog classifier (pure, position-independent)."""

    def test_exact_alias_unknown(self):
        self.assertEqual(classify_kp_label("MSM").status, "exact")
        self.assertEqual(classify_kp_label("GENERAL POP.").status, "exact")
        self.assertEqual(classify_kp_label("General population"),
                         ("alias", "GENERAL POP."))
        self.assertEqual(classify_kp_label("PWIDs"), ("alias", "PWID"))
        self.assertEqual(classify_kp_label("Martians").status, "unknown")

    def test_pwud_not_aliased_to_pwid(self):
        # PWUD is a distinct recognised category, never folded into PWID.
        self.assertEqual(classify_kp_label("PWUD"), ("exact", "PWUD"))

    def test_indicator_configured_value_widens_recognised_set(self):
        # A value the indicator itself declares is accepted even if off-catalog.
        self.assertEqual(
            classify_kp_label("Transgender", allowed_values=["Transgender"]),
            ("exact", "Transgender"),
        )
        self.assertEqual(classify_kp_label("Transgender").status, "unknown")


class KpParseTest(TestCase):
    """Parse-level behaviour of the smart importer (no HTTP)."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user(
            username="kp_admin", email="kp_admin@example.com", password="x",
            role="admin",
        )
        cls.project = Project.objects.create(
            name="P", code="KPV-1", start_date=date(2025, 4, 1),
            end_date=date(2026, 3, 31), created_by=cls.admin,
        )
        cls.org = Organization.objects.create(name="Org", code="KPORG", type="cso")
        cls.ind = Indicator.objects.create(
            name="People reached (KP)", code="KPV_IND", type="number",
            category="hiv_prevention", aggregate_disaggregation_config=KP_CONFIG,
            created_by=cls.admin,
        )

    def _parse(self, cells, **kw):
        data = make_workbook(cells, indicator_id=self.ind.id,
                             project_id=self.project.id, org_id=self.org.id, **kw)
        return rw.parse_workbook(BytesIO(data))

    def test_v2_current_arrangement_no_warning(self):
        parsed = self._parse([
            ("GENERAL POP.", "Male", "10-14", 3),
            ("GENERAL POP.", "Female", "15-19", 4),
            ("PWID", "Male", "10-14", 5),
        ], kp_arrangement="v2")
        self.assertEqual(parsed.warnings, [])
        self.assertEqual(len(parsed.indicators), 1)
        d = parsed.indicators[0].value["disaggregates"]
        self.assertEqual(d["GENERAL POP."]["Male"]["10-14"], 3)
        self.assertEqual(d["GENERAL POP."]["Female"]["15-19"], 4)
        self.assertEqual(d["PWID"]["Male"]["10-14"], 5)
        self.assertEqual(parsed.indicators[0].value["total"], 12)

    def test_v1_old_arrangement_imports_with_notice(self):
        # Old order + old spelling; v1 marker. Must map forward + warn.
        parsed = self._parse([
            ("General population", "Male", "10-14", 3),
            ("PWIDs", "Female", "15-19", 2),
        ], kp_arrangement="v1")
        self.assertIn(rw.OLD_KP_TEMPLATE_MESSAGE, parsed.warnings)
        d = parsed.indicators[0].value["disaggregates"]
        self.assertEqual(d["GENERAL POP."]["Male"]["10-14"], 3)
        self.assertEqual(d["PWID"]["Female"]["15-19"], 2)
        self.assertNotIn("General population", d)
        self.assertNotIn("PWIDs", d)

    def test_v1_and_v2_yield_identical_internal_categories(self):
        # Same data per category, different column order + spelling → identical.
        v2 = self._parse([
            ("MSM", "Male", "10-14", 1),
            ("FSW", "Female", "10-14", 2),
            ("GENERAL POP.", "Male", "15-19", 3),
        ], kp_arrangement="v2").indicators[0].value
        v1 = self._parse([
            ("General population", "Male", "15-19", 3),
            ("FSW", "Female", "10-14", 2),
            ("MSM", "Male", "10-14", 1),
        ], kp_arrangement="v1").indicators[0].value
        self.assertEqual(v1["disaggregates"], v2["disaggregates"])
        self.assertEqual(v1["total"], v2["total"])

    def test_missing_marker_but_current_labels_no_noise(self):
        # A pre-feature workbook (no marker) whose labels are all current must
        # NOT raise a spurious "older template" notice.
        parsed = self._parse([
            ("GENERAL POP.", "Male", "10-14", 3),
        ], kp_arrangement=None)
        self.assertEqual(parsed.warnings, [])

    def test_unknown_category_flagged_and_not_stored(self):
        parsed = self._parse([
            ("GENERAL POP.", "Male", "10-14", 3),
            ("Martians", "Male", "10-14", 9),
        ], kp_arrangement="v2")
        item = parsed.indicators[0]
        self.assertTrue(item.problems)
        self.assertIn("Martians", item.problems[0])
        # The unknown category is NOT present under any label.
        d = item.value.get("disaggregates", {})
        self.assertNotIn("Martians", d)

    def test_duplicate_conflicting_labels_merged_and_reported(self):
        # Two spellings of one category, same (sex, band) → summed + conflict note.
        parsed = self._parse([
            ("GENERAL POP.", "Male", "10-14", 4),
            ("General population", "Male", "10-14", 6),
        ], kp_arrangement="v1")
        d = parsed.indicators[0].value["disaggregates"]
        self.assertEqual(d["GENERAL POP."]["Male"]["10-14"], 10)  # 4 + 6
        self.assertTrue(any("combined into" in w for w in parsed.warnings))

    def test_age_and_sex_stay_bound_to_their_kp(self):
        parsed = self._parse([
            ("PWID", "Male", "10-14", 1),
            ("PWID", "Female", "15-19", 2),
            ("GENERAL POP.", "Male", "15-19", 7),
        ], kp_arrangement="v2")
        d = parsed.indicators[0].value["disaggregates"]
        self.assertEqual(d["PWID"], {"Male": {"10-14": 1}, "Female": {"15-19": 2}})
        self.assertEqual(d["GENERAL POP."], {"Male": {"15-19": 7}})

    def test_total_correct_after_mapping(self):
        parsed = self._parse([
            ("General population", "Male", "10-14", 4),
            ("GENERAL POP.", "Male", "10-14", 6),
            ("PWIDs", "Female", "15-19", 5),
        ], kp_arrangement="v1")
        self.assertEqual(parsed.indicators[0].value["total"], 15)  # 4+6+5

    def test_missing_required_metadata_rejected(self):
        # Drop project_id from a hand-built workbook → friendly rejection.
        data = make_workbook(
            [("GENERAL POP.", "Male", "10-14", 1)],
            indicator_id=self.ind.id, project_id="", org_id=self.org.id,
        )
        with self.assertRaises(rw.WorkbookError):
            rw.parse_workbook(BytesIO(data))

    def test_non_kp_primary_untouched(self):
        # An indicator whose primary axis is NOT key-population must pass through
        # unchanged — no KP validation, no false "unknown" flagging.
        other = Indicator.objects.create(
            name="Posts by platform", code="KPV_PLATFORM", type="number",
            category="hiv_prevention",
            aggregate_disaggregation_config={
                "enabled": True, "layout": "list",
                "dimensions": [{"key": "social_media_platform",
                                "label": "Social Media Platform",
                                "values": ["Facebook", "WhatsApp"]}],
            },
            created_by=self.admin,
        )
        data = make_workbook(
            [("Facebook", "All", "Value", 8), ("WhatsApp", "All", "Value", 2)],
            indicator_id=other.id, project_id=self.project.id, org_id=self.org.id,
        )
        parsed = rw.parse_workbook(BytesIO(data))
        self.assertEqual(parsed.warnings, [])
        d = parsed.indicators[0].value["disaggregates"]
        self.assertIn("Facebook", d)
        self.assertIn("WhatsApp", d)


class KpImportHttpTest(APITestCase):
    """End-to-end import path: warnings surfaced, unknown rejected, canonical
    stored, historical data untouched."""

    @classmethod
    def setUpTestData(cls):
        cls.org = Organization.objects.create(name="Home of Hope", code="KPH_HOH", type="cso")
        cls.admin = User.objects.create_user(
            username="kph_admin", email="kph_admin@example.com", password="x",
            role="admin", organization=cls.org,
        )
        cls.officer = User.objects.create_user(
            username="kph_off", email="kph_off@example.com", password="x",
            role="officer", organization=cls.org,
        )
        cls.project = Project.objects.create(
            name="Sesigo", code="KPH-1", start_date=date(2025, 4, 1),
            end_date=date(2026, 3, 31), created_by=cls.admin,
        )
        cls.project.organizations.add(cls.org)
        cls.ind = Indicator.objects.create(
            name="People reached (KP)", code="KPH_IND", type="number",
            category="hiv_prevention", aggregate_disaggregation_config=KP_CONFIG,
            created_by=cls.admin,
        )
        pi = ProjectIndicator.objects.create(project=cls.project, indicator=cls.ind, q3_target=100)
        ProjectIndicatorOrganizationTarget.objects.create(
            project_indicator=pi, organization=cls.org, q3_target=100,
        )
        cls.ind.organizations.add(cls.org)

    IMPORT_URL = "/api/aggregates/import-reporting-workbook/"

    def _post(self, cells, **kw):
        data = make_workbook(cells, indicator_id=self.ind.id,
                             project_id=self.project.id, org_id=self.org.id, **kw)
        self.client.force_authenticate(self.officer)
        upload = BytesIO(data)
        upload.name = "filled.xlsx"
        return self.client.post(self.IMPORT_URL, {"file": upload}, format="multipart")

    def test_old_template_imports_canonicalised_with_warning(self):
        resp = self._post([
            ("General population", "Male", "10-14", 3),
            ("PWIDs", "Female", "15-19", 2),
        ], kp_arrangement="v1")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)
        self.assertIn(rw.OLD_KP_TEMPLATE_MESSAGE, resp.data.get("warnings", []))
        agg = Aggregate.objects.get(indicator=self.ind, project=self.project, organization=self.org)
        d = agg.value["disaggregates"]
        # Stored under canonical labels only (pre_save guard + parser agree).
        self.assertIn("GENERAL POP.", d)
        self.assertIn("PWID", d)
        self.assertNotIn("General population", d)
        self.assertNotIn("PWIDs", d)

    def test_unknown_category_rejected_nothing_written(self):
        resp = self._post([
            ("Martians", "Male", "10-14", 9),
        ], kp_arrangement="v2")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST, resp.data)
        joined = " ".join(e.get("error", "") for e in resp.data.get("errors", []))
        self.assertIn("Martians", joined)
        self.assertFalse(
            Aggregate.objects.filter(indicator=self.ind, organization=self.org).exists()
        )

    def test_import_does_not_touch_unrelated_historical_aggregate(self):
        # A pre-existing aggregate for a different period must be byte-identical
        # after importing a current-period report.
        historical = Aggregate.objects.create(
            indicator=self.ind, project=self.project, organization=self.org,
            period_start=date(2025, 7, 1), period_end=date(2025, 9, 30),
            value={"disaggregates": {"GENERAL POP.": {"Male": {"10-14": 99}}}, "total": 99},
        )
        before = dict(historical.value)
        resp = self._post([("GENERAL POP.", "Male", "10-14", 1)], kp_arrangement="v2")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)
        historical.refresh_from_db()
        self.assertEqual(historical.value, before)
